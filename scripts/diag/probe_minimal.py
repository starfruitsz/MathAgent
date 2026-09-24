"""诊断：隔离测试——最小 docx（1 个公式）里的 OLE 对象能否被 Word 激活。

逐步定位是哪一步破坏了对象：
  ① 只做 MathType 转换（不重绘预览图）
  ② 再做预览图重绘 + 写回

用法：
    python scripts/diag/probe_minimal.py
"""

from __future__ import annotations

import subprocess
import sys
import time
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from report import equations as EQ  # noqa: E402

WORK = ROOT / "paper" / "_probe_min"


def build_omml() -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_paragraph("测试：")
    p = doc.add_paragraph()
    EQ.append_omml(p, r"E_{gij}^{hor}(q) = (d_{ij} / L_{g}(q)) \cdot E_{g}^{use}",
                   display=False)
    WORK.mkdir(parents=True, exist_ok=True)
    src = WORK / "omml.docx"
    doc.save(str(src))
    return src


def activate_and_report(docx: Path, label: str) -> bool:
    import win32com.client as win32

    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    ok = False
    try:
        doc = word.Documents.Open(str(docx))
        try:
            time.sleep(2.0)
            n = doc.InlineShapes.Count
            types = {}
            for i in range(1, n + 1):
                t = int(doc.InlineShapes.Item(i).Type)
                types[t] = types.get(t, 0) + 1
            print(f"  [{label}] InlineShapes={n} 类型分布={types}")
            oles = [i for i in range(1, n + 1)
                    if int(doc.InlineShapes.Item(i).Type) == 1]
            if not oles:
                print(f"  [{label}] ❌ 无 OLE 对象")
                return False
            sh = doc.InlineShapes.Item(oles[0])
            try:
                prog = sh.OLEFormat.ProgID
                print(f"  [{label}] ProgID = {prog}")
            except Exception as exc:  # noqa: BLE001
                print(f"  [{label}] ❌ OLEFormat.ProgID 读取失败：{str(exc)[:120]}")
                return False
            try:
                sh.OLEFormat.Activate()
                time.sleep(4.0)
                print(f"  [{label}] ✅ Activate 成功（对象可编辑）")
                ok = True
            except Exception as exc:  # noqa: BLE001
                print(f"  [{label}] ❌ Activate 失败：{str(exc)[:160]}")
        finally:
            try:
                doc.Close(SaveChanges=False)
            except Exception:  # noqa: BLE001
                pass
    finally:
        try:
            word.Quit()
        except Exception:  # noqa: BLE001
            pass
        subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"],
                       capture_output=True, check=False)
        time.sleep(1.0)
    return ok


def main() -> int:
    WORK.mkdir(parents=True, exist_ok=True)
    src = build_omml()
    print(f"OMML 源：{src.name}")

    # ① 仅转换，不重绘预览图
    a = WORK / "step1_convert_only.docx"
    for p in (a,):
        p.unlink(missing_ok=True)
    from docx_equation import convert_omml_docx_to_mathtype
    from docx_equation.shared import mathml as _m

    br = EQ.find_browser()
    if br:
        _m._find_chrome = lambda *_a, **_k: Path(br)  # noqa: SLF001
    work = WORK / "work1"
    if work.exists():
        import shutil
        shutil.rmtree(work)
    n = convert_omml_docx_to_mathtype(
        str(src), str(a), work_dir=str(work),
        omml2mathml_xsl=str(EQ.WIN_OMML_XSL),
        preview_pt_per_px=EQ.PREVIEW_PT_PER_PX, mathtype_version="DSMT4")
    print(f"① 转换完成（{n} 个公式），未重绘预览图 → {a.name}")
    ok1 = activate_and_report(a, "① 仅转换")

    # ② 追加预览图重绘 + 写回
    b = WORK / "step2_reembed.docx"
    import shutil
    shutil.copy(a, b)
    EQ.render_previews(work / "mathml", work / "preview_png", browser=br)
    EQ.reembed_previews(b, work / "preview_png")
    print(f"② 已重绘预览图并写回 → {b.name}")
    ok2 = activate_and_report(b, "② 重绘后")

    print(f"\n结论：仅转换可激活={ok1}；重绘后仍可激活={ok2}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
